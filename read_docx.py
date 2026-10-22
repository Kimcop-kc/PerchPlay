from docx import Document
import sys, io, json

# python-docx returns text as Unicode strings.
# Force stdout to UTF-8 so pwsh/ASCII pipe won't garble CJK.
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = Document(r'C:\Users\86177\Desktop\PerchPlay 播放器bug修复任务书.docx')
out = []
for i, p in enumerate(doc.paragraphs, 1):
    style = p.style.name if p.style else 'None'
    text = p.text.strip()
    if text:
        out.append(f'[P{i}|{style}] {text}')
for ti, table in enumerate(doc.tables):
    out.append(f'\n=== Table {ti+1} ===')
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        out.append(f'  R{ri}: ' + ' | '.join(cells))

# Write to UTF-8 file
with open(r'C:\Users\86177\Documents\deepseek-harness\default-workspace\PerchPlay\docx_output.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('DONE', len(out), 'lines')